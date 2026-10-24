import io
import json
import os
import sys
import traceback
from datetime import datetime

import tkinter as tk
from tkinter import filedialog, messagebox

import customtkinter as ctk
from PIL import Image, ImageDraw, ImageFont, ImageGrab, ImageOps, ImageTk
from docx import Document
from docx.shared import Inches, Mm, Pt

FONT = "Segoe UI"
CJK_FONT = "Microsoft YaHei UI"
DEFAULT_LANGUAGE = "en"

LIGHT = {
    "bg": "#F2F5FA",
    "panel": "#F8FAFD",
    "card": "#FFFFFF",
    "card_hover": "#F4F8FF",
    "selected_bg": "#F1F6FF",
    "border": "#E3E9F2",
    "text": "#141A23",
    "muted": "#687385",
    "accent": "#2F6BFF",
    "accent_hover": "#2459E0",
    "accent_soft": "#EAF1FF",
    "green": "#17A35A",
    "green_hover": "#128A4B",
    "danger": "#D92D20",
    "danger_hover": "#FEF3F2",
    "ghost_hover": "#EEF3FA",
    "files_hover": "#F1F5FB",
    "scrollbar": "#D6DEEA",
    "scrollbar_hover": "#BAC7D9",
    "tile_bg": "#F1F5FA",
    "tile_border": "#E3E9F2",
    "toast_info": ("#EAF1FF", "#2F6BFF", "#D6E4FF"),
    "toast_success": ("#E9F9F0", "#0B7A45", "#CFEEDF"),
    "toast_error": ("#FEF3F2", "#D92D20", "#FADBD7"),
}

DARK = {
    "bg": "#0F141C",
    "panel": "#151C27",
    "card": "#1B2432",
    "card_hover": "#233047",
    "selected_bg": "#1A2B47",
    "border": "#2A3648",
    "text": "#E8EEF8",
    "muted": "#9AA7BA",
    "accent": "#6A96FF",
    "accent_hover": "#82A8FF",
    "accent_soft": "#1D3150",
    "green": "#22B573",
    "green_hover": "#2ECC82",
    "danger": "#FF7168",
    "danger_hover": "#3A2022",
    "ghost_hover": "#202B3C",
    "files_hover": "#243044",
    "scrollbar": "#334154",
    "scrollbar_hover": "#45566D",
    "tile_bg": "#202A39",
    "tile_border": "#2C394C",
    "toast_info": ("#1D3150", "#9CBBFF", "#2E4A73"),
    "toast_success": ("#123326", "#55D99A", "#1E5340"),
    "toast_error": ("#3A2022", "#FF9A92", "#5A2B2C"),
}


def _pair(name):
    return (LIGHT[name], DARK[name])


BG = _pair("bg")
PANEL = _pair("panel")
CARD = _pair("card")
CARD_HOVER = _pair("card_hover")
SELECTED_BG = _pair("selected_bg")
BORDER = _pair("border")
TEXT = _pair("text")
MUTED = _pair("muted")
ACCENT = _pair("accent")
ACCENT_HOVER = _pair("accent_hover")
ACCENT_SOFT = _pair("accent_soft")
GREEN = _pair("green")
GREEN_HOVER = _pair("green_hover")
DANGER = _pair("danger")
DANGER_HOVER = _pair("danger_hover")
GHOST_HOVER = _pair("ghost_hover")
FILES_HOVER = _pair("files_hover")
SCROLLBAR = _pair("scrollbar")
SCROLLBAR_HOVER = _pair("scrollbar_hover")

GLYPH_FONT = r"C:\Windows\Fonts\segmdl2.ttf"

IMAGE_EXTENSIONS = (".png", ".jpg", ".jpeg", ".bmp", ".gif", ".webp", ".tif", ".tiff")

SETTINGS_DIR = os.path.join(
    os.environ.get("APPDATA") or os.path.expanduser("~"), "PhotoToWord")
SETTINGS_PATH = os.path.join(SETTINGS_DIR, "settings.json")


def load_settings():
    try:
        with open(SETTINGS_PATH, encoding="utf-8-sig") as file:
            data = json.load(file)
        return data if isinstance(data, dict) else {}
    except (OSError, ValueError):
        return {}


def save_settings(settings):
    try:
        os.makedirs(SETTINGS_DIR, exist_ok=True)
        temp_path = SETTINGS_PATH + ".tmp"
        with open(temp_path, "w", encoding="utf-8") as file:
            json.dump(settings, file, ensure_ascii=False, indent=2)
        os.replace(temp_path, SETTINGS_PATH)
    except OSError:
        pass


def load_theme():
    mode = load_settings().get("appearance_mode")
    return mode if mode in ("light", "dark") else "light"


def save_theme(mode):
    settings = load_settings()
    settings["appearance_mode"] = mode
    save_settings(settings)


def save_language(code):
    settings = load_settings()
    settings["language"] = code
    save_settings(settings)


def _resource_base():
    if getattr(sys, "frozen", False):
        return getattr(sys, "_MEIPASS", os.path.dirname(sys.executable))
    return os.path.dirname(os.path.abspath(__file__))


LOCALES_DIR = os.path.join(_resource_base(), "locales")
_LOCALE_CACHE = {}


def _read_json(path):
    try:
        with open(path, encoding="utf-8-sig") as file:
            data = json.load(file)
        return data if isinstance(data, dict) else {}
    except (OSError, ValueError):
        return {}


def load_locale(code):
    if code not in _LOCALE_CACHE:
        _LOCALE_CACHE[code] = _read_json(
            os.path.join(LOCALES_DIR, f"{code}.json"))
    return _LOCALE_CACHE[code]


def available_language_codes():
    preferred = ("en", "ru", "zh")
    try:
        names = os.listdir(LOCALES_DIR)
    except OSError:
        names = []
    codes = []
    for name in names:
        if not name.endswith(".json") or name == "index.json":
            continue
        code = name[:-5]
        if load_locale(code).get("_meta"):
            codes.append(code)
    order = {code: index for index, code in enumerate(preferred)}
    codes.sort(key=lambda code: (order.get(code, len(preferred)), code))
    return codes


def available_languages():
    return [(code, load_locale(code).get("_meta", {}))
            for code in available_language_codes()]


def load_language():
    code = load_settings().get("language")
    return code if code in available_language_codes() else None


def load_index():
    return _read_json(os.path.join(LOCALES_DIR, "index.json"))


class Translator:
    def __init__(self, code):
        codes = available_language_codes()
        self.code = code if code in codes else DEFAULT_LANGUAGE
        self.data = load_locale(self.code)
        self.fallback = load_locale(DEFAULT_LANGUAGE)
        self.meta = self.data.get("_meta", {})

    def __call__(self, key, **values):
        text = self.data.get(key)
        if not isinstance(text, str):
            text = self.fallback.get(key)
        if not isinstance(text, str):
            text = key
        if not values:
            return text
        try:
            return text.format(**values)
        except (KeyError, ValueError):
            return text


# Icon glyphs from Segoe MDL2 Assets.
GLYPHS = {
    "paste": "\uE77F",
    "files": "\uE8B7",
    "left": "\uE76B",
    "right": "\uE76C",
    "delete": "\uE74D",
    "save": "\uE74E",
    "clear": "\uE75C",
    "photo": "\uEB9F",
    "doc": "\uE8A5",
    "sun": "\uE706",
    "moon": "\uE708",
    "globe": "\uE774",
}


def desktop_path():
    """Return the Desktop folder, including OneDrive-redirected setups."""
    if sys.platform == "win32":
        try:
            import winreg

            key_path = r"Software\Microsoft\Windows\CurrentVersion\Explorer\User Shell Folders"
            with winreg.OpenKey(winreg.HKEY_CURRENT_USER, key_path) as key:
                value, _ = winreg.QueryValueEx(key, "Desktop")
            path = os.path.expandvars(value)
            if os.path.isdir(path):
                return path
        except Exception:
            pass
    return os.path.join(os.path.expanduser("~"), "Desktop")


def _to_rgb(image):
    if image.mode == "RGB":
        return image
    if image.mode in ("RGBA", "LA") or (image.mode == "P" and "transparency" in image.info):
        rgba = image.convert("RGBA")
        background = Image.new("RGB", rgba.size, "white")
        background.paste(rgba, mask=rgba.split()[-1])
        return background
    return image.convert("RGB")


def _size_inches(image):
    dpi = image.info.get("dpi") or (96, 96)
    try:
        dx = float(dpi[0]) or 96.0
    except Exception:
        dx = 96.0
    try:
        dy = float(dpi[1]) or 96.0
    except Exception:
        dy = 96.0
    return image.width / dx, image.height / dy


def save_docx(images, path):
    """Save images to Word: a number, the photo below it, and so on."""
    document = Document()

    section = document.sections[0]
    section.page_width = Mm(210)   # A4
    section.page_height = Mm(297)
    section.left_margin = section.right_margin = Inches(0.7)
    section.top_margin = section.bottom_margin = Inches(0.7)

    emu_per_inch = 914400
    usable_w = (section.page_width - section.left_margin - section.right_margin) / emu_per_inch
    usable_h = (section.page_height - section.top_margin - section.bottom_margin) / emu_per_inch - 0.5

    for number, image in enumerate(images, 1):
        paragraph = document.add_paragraph()
        run = paragraph.add_run(f"{number}.")
        run.bold = True
        run.font.size = Pt(14)
        paragraph.paragraph_format.space_before = Pt(12 if number > 1 else 0)
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.paragraph_format.keep_with_next = True

        width_in, height_in = _size_inches(image)
        scale = min(1.0, usable_w / width_in, usable_h / height_in)
        width_in = max(0.6, width_in * scale)

        stream = io.BytesIO()
        _to_rgb(image).save(stream, format="PNG")
        stream.seek(0)

        document.add_picture(stream, width=Inches(width_in))
        document.paragraphs[-1].paragraph_format.space_after = Pt(8)

    document.save(path)


# ---------- UI graphics ----------

def _glyph_image(char, color, size):
    """Render a Segoe MDL2 Assets icon at high resolution."""
    scale = 4
    pixels = size * scale
    image = Image.new("RGBA", (pixels, pixels), (0, 0, 0, 0))
    draw = ImageDraw.Draw(image)
    try:
        font = ImageFont.truetype(GLYPH_FONT, int(pixels * 0.82))
    except OSError:
        font = ImageFont.truetype(r"C:\Windows\Fonts\segoeui.ttf", int(pixels * 0.72))
    draw.text((pixels / 2, pixels / 2), char, font=font, fill=color, anchor="mm")
    return image


def _split_color(color):
    if isinstance(color, (tuple, list)) and len(color) == 2:
        return color[0], color[1]
    return color, color


def _icon(char, color, size=16):
    light_color, dark_color = _split_color(color)
    return ctk.CTkImage(
        light_image=_glyph_image(char, light_color, size),
        dark_image=_glyph_image(char, dark_color, size),
        size=(size, size),
    )


def _logo_image(char, fg, bg, size, radius_ratio=0.28):
    scale = 4
    pixels = size * scale
    image = Image.new("RGBA", (pixels, pixels), (0, 0, 0, 0))
    draw = ImageDraw.Draw(image)
    draw.rounded_rectangle([0, 0, pixels - 1, pixels - 1],
                           radius=int(pixels * radius_ratio), fill=bg)
    try:
        font = ImageFont.truetype(GLYPH_FONT, int(pixels * 0.5))
    except OSError:
        font = ImageFont.truetype(r"C:\Windows\Fonts\segoeui.ttf", int(pixels * 0.5))
    draw.text((pixels / 2, pixels / 2 + pixels * 0.01), char, font=font,
              fill=fg, anchor="mm")
    return image


def _round_icon(char, fg, bg, size=104):
    scale = 2
    pixels = size * scale
    image = Image.new("RGBA", (pixels, pixels), (0, 0, 0, 0))
    draw = ImageDraw.Draw(image)
    draw.ellipse([0, 0, pixels - 1, pixels - 1], fill=bg)
    try:
        font = ImageFont.truetype(GLYPH_FONT, int(pixels * 0.4))
    except OSError:
        font = ImageFont.truetype(r"C:\Windows\Fonts\segoeui.ttf", int(pixels * 0.4))
    draw.text((pixels / 2, pixels / 2), char, font=font, fill=fg, anchor="mm")
    return image


def _thumb_tile(image, size, radius=12, padding=8,
                bg="#F1F5FA", border="#E3E9F2"):
    """Build a square rounded tile with a preview of the photo."""
    scale = 2
    pixels = size * scale
    radius *= scale
    padding *= scale

    tile = Image.new("RGBA", (pixels, pixels), (0, 0, 0, 0))
    draw = ImageDraw.Draw(tile)
    draw.rounded_rectangle([0, 0, pixels - 1, pixels - 1],
                           radius=radius, fill=bg,
                           outline=border, width=scale)

    preview = _to_rgb(image).convert("RGBA")
    preview.thumbnail((pixels - padding * 2, pixels - padding * 2), Image.LANCZOS)
    tile.paste(preview, ((pixels - preview.width) // 2, (pixels - preview.height) // 2))

    mask = Image.new("L", (pixels, pixels), 0)
    ImageDraw.Draw(mask).rounded_rectangle([0, 0, pixels - 1, pixels - 1],
                                           radius=radius, fill=255)
    tile.putalpha(mask)
    return tile


def _set_dark_titlebar(window, dark):
    """Match the native Windows title bar to the current theme."""
    if sys.platform != "win32":
        return
    try:
        import ctypes

        hwnd = ctypes.windll.user32.GetParent(window.winfo_id())
        if not hwnd:
            hwnd = window.winfo_id()
        value = ctypes.c_int(1 if dark else 0)
        for attribute in (20, 19):
            result = ctypes.windll.dwmapi.DwmSetWindowAttribute(
                hwnd, attribute, ctypes.byref(value), ctypes.sizeof(value))
            if result == 0:
                break
    except Exception:
        pass


class LanguageDialog(ctk.CTkToplevel):
    WIDTH = 480
    HEIGHT = 412

    def __init__(self, parent, current=None, localized_title=None,
                 localized_subtitle=None):
        super().__init__(parent)
        self.parent = parent
        self.result = None
        self._current = current
        self._index = load_index()
        self._title_text = localized_title or self._index.get("title") or "Language"
        self._subtitle_text = localized_subtitle or self._index.get("subtitle") or ""

        self.title(self._title_text)
        self.configure(fg_color=BG)
        self.resizable(False, False)
        self.protocol("WM_DELETE_WINDOW", self._cancel)
        if parent.winfo_viewable():
            self.transient(parent)

        self._build()
        self.after(30, self._show)

    def _build(self):
        inner = ctk.CTkFrame(self, fg_color="transparent")
        inner.pack(fill="both", expand=True, padx=26, pady=22)

        ctk.CTkLabel(
            inner, text=self._title_text, font=ctk.CTkFont(FONT, 20, "bold"),
            text_color=TEXT, anchor="w", justify="left",
        ).pack(fill="x")
        if self._subtitle_text:
            ctk.CTkLabel(
                inner, text=self._subtitle_text, font=ctk.CTkFont(FONT, 11),
                text_color=MUTED, anchor="w", justify="left", wraplength=420,
            ).pack(fill="x", pady=(4, 16))

        for code, meta in available_languages():
            selected = code == self._current
            name = meta.get("name", code)
            native = meta.get("native", "")
            label = name
            if native and native != name:
                label = f"{name}  ·  {native}"
            ctk.CTkButton(
                inner, text=label, anchor="w", height=50,
                font=ctk.CTkFont(FONT, 13, "bold"),
                fg_color=SELECTED_BG if selected else CARD,
                hover_color=CARD_HOVER,
                text_color=TEXT,
                border_width=2 if selected else 1,
                border_color=ACCENT if selected else BORDER,
                corner_radius=12,
                command=lambda value=code: self._choose(value),
            ).pack(fill="x", pady=5)

        hint = self._index.get("hint", "")
        if hint and not self._title_text == self._index.get("title"):
            hint = ""
        if hint:
            ctk.CTkLabel(
                inner, text=hint, font=ctk.CTkFont(FONT, 10),
                text_color=MUTED, wraplength=420, justify="center",
            ).pack(fill="x", pady=(12, 0))

    def _show(self):
        self.update_idletasks()
        screen_w = self.winfo_screenwidth()
        screen_h = self.winfo_screenheight()
        x = max(0, (screen_w - self.WIDTH) // 2)
        y = max(0, (screen_h - self.HEIGHT) // 2)
        self.geometry(f"{self.WIDTH}x{self.HEIGHT}+{x}+{y}")
        self.lift()
        self.focus_force()
        try:
            self.grab_set()
        except tk.TclError:
            pass
        _set_dark_titlebar(self, ctk.get_appearance_mode().lower() == "dark")

    def _choose(self, code):
        self.result = code
        self.destroy()

    def _cancel(self):
        self.result = None
        self.destroy()


class App:
    THUMB = 150
    CARD_W = 180
    CARD_H = 214

    def __init__(self, root, language):
        self.root = root
        self.lang = language
        self.t = Translator(language)
        self.meta = self.t.meta
        self.font_family = CJK_FONT if self.lang == "zh" else FONT
        self.items = []
        self.selected = None
        self._cols = 0
        self._toast_job = None
        self._showing_empty = None

        self.theme = load_theme()
        ctk.set_appearance_mode(self.theme)

        screen_w = root.winfo_screenwidth()
        screen_h = root.winfo_screenheight()

        try:
            scale = float(ctk.ScalingTracker.get_window_scaling(root))
        except Exception:
            scale = 1.0
        if not 0.5 <= scale <= 3.0:
            scale = 1.0

        width = min(int(1140 * scale), screen_w - 80)
        height = min(int(800 * scale), screen_h - 110)
        min_width = min(int(980 * scale), width)
        min_height = min(int(620 * scale), height)
        x = max(0, (screen_w - width) // 2)
        y = max(0, (screen_h - height) // 2)

        root.title(self.t("app_title"))
        root.geometry(f"{width}x{height}+{x}+{y}")
        root.minsize(min_width, min_height)
        root.configure(fg_color=BG)
        try:
            icon = _logo_image(GLYPHS["photo"], "white", LIGHT["accent"], 64,
                               radius_ratio=0.24)
            self._window_icon = ImageTk.PhotoImage(icon)
            root.iconphoto(True, self._window_icon)
        except Exception:
            pass

        self._build_fonts()
        self._build_ui()
        self._update_theme_button()
        self._build_empty()
        self._bind_keys()
        self._render()
        _set_dark_titlebar(root, self.theme == "dark")

    # ---------- fonts ----------

    def _build_fonts(self):
        family = self.font_family
        self.f_title = ctk.CTkFont(family, 24, "bold")
        self.f_subtitle = ctk.CTkFont(family, 11)
        self.f_primary = ctk.CTkFont(family, 12, "bold")
        self.f_secondary = ctk.CTkFont(family, 11)
        self.f_ghost = ctk.CTkFont(family, 11)
        self.f_badge = ctk.CTkFont(family, 10, "bold")
        self.f_card = ctk.CTkFont(family, 11)
        self.f_count = ctk.CTkFont(family, 14, "bold")
        self.f_hint = ctk.CTkFont(family, 10)
        self.f_empty_title = ctk.CTkFont(family, 20, "bold")
        self.f_empty_sub = ctk.CTkFont(family, 12)
        self.f_chip = ctk.CTkFont(family, 10, "bold")

    # ---------- interface ----------

    def _build_ui(self):
        header = ctk.CTkFrame(self.root, fg_color="transparent")
        header.pack(fill="x", padx=24, pady=(20, 12))

        logo = ctk.CTkLabel(
            header, text="", width=46, height=46,
            image=ctk.CTkImage(
                light_image=_logo_image(GLYPHS["photo"], "white", LIGHT["accent"], 46),
                dark_image=_logo_image(GLYPHS["photo"], "white", DARK["accent"], 46),
                size=(46, 46),
            ),
        )
        logo.pack(side="left")

        titles = ctk.CTkFrame(header, fg_color="transparent")
        titles.pack(side="left", padx=(12, 0))
        ctk.CTkLabel(titles, text=self.t("app_title"), font=self.f_title,
                     text_color=TEXT).pack(anchor="w")
        ctk.CTkLabel(titles, text=self.t("app_subtitle"),
                     font=self.f_subtitle, text_color=MUTED).pack(anchor="w")

        self._globe_icon = _icon(GLYPHS["globe"], MUTED, 14)
        self.lang_button = ctk.CTkButton(
            header, text=self.meta.get("short", "EN"), image=self._globe_icon,
            compound="left", command=self.open_language_dialog,
            font=self.f_badge, fg_color=CARD, hover_color=GHOST_HOVER,
            text_color=MUTED, border_width=1, border_color=BORDER,
            corner_radius=14, height=28, width=72,
        )
        self.lang_button.pack(side="right", padx=(0, 8), pady=8)

        self._moon_icon = _icon(GLYPHS["moon"], MUTED, 14)
        self._sun_icon = _icon(GLYPHS["sun"], MUTED, 14)
        self.theme_button = ctk.CTkButton(
            header, text=self.t("theme_dark"), image=self._moon_icon, compound="left",
            command=self.toggle_theme, font=self.f_badge,
            fg_color=CARD, hover_color=GHOST_HOVER, text_color=MUTED,
            border_width=1, border_color=BORDER, corner_radius=14,
            height=28, width=108,
        )
        self.theme_button.pack(side="right", padx=(0, 8), pady=8)

        badge = ctk.CTkLabel(header, text=self.t("docx_badge"), font=self.f_badge,
                             text_color=MUTED, fg_color=CARD,
                             corner_radius=13, width=92, height=28,
                             border_width=1, border_color=BORDER)
        badge.pack(side="right", pady=8)

        toolbar = ctk.CTkFrame(self.root, fg_color="transparent")
        toolbar.pack(fill="x", padx=24)

        self.paste_button = ctk.CTkButton(
            toolbar, text=self.t("paste"), image=_icon(GLYPHS["paste"], "white", 16),
            compound="left", command=self.paste,
            font=self.f_primary, fg_color=ACCENT, hover_color=ACCENT_HOVER,
            corner_radius=12, height=42, width=196,
        )
        self.paste_button.pack(side="left")

        self.files_button = ctk.CTkButton(
            toolbar, text=self.t("add_files"), image=_icon(GLYPHS["files"], MUTED, 16),
            compound="left", command=self.add_files,
            font=self.f_secondary, fg_color=CARD, hover_color=FILES_HOVER,
            text_color=TEXT, border_width=1, border_color=BORDER,
            corner_radius=12, height=42, width=164,
        )
        self.files_button.pack(side="left", padx=(8, 0))

        ctk.CTkFrame(toolbar, width=1, height=24, fg_color=BORDER).pack(
            side="left", padx=14, pady=9)

        self._ghost_button(toolbar, self.t("move_earlier"), GLYPHS["left"], MUTED,
                           lambda: self.move_selected(-1)).pack(side="left")
        self._ghost_button(toolbar, self.t("move_later"), GLYPHS["right"], MUTED,
                           lambda: self.move_selected(1)).pack(side="left", padx=(4, 0))
        self._ghost_button(toolbar, self.t("delete"), GLYPHS["delete"], DANGER,
                           self.delete_selected,
                           hover=DANGER_HOVER).pack(side="left", padx=(12, 0))
        self._ghost_button(toolbar, self.t("clear"), GLYPHS["clear"], MUTED,
                           self.clear_all).pack(side="left", padx=(4, 0))

        panel = ctk.CTkFrame(self.root, fg_color=PANEL, corner_radius=20,
                             border_width=1, border_color=BORDER)
        panel.pack(fill="both", expand=True, padx=24, pady=(16, 12))
        self.panel = panel

        self.cards = ctk.CTkScrollableFrame(
            panel, fg_color="transparent", corner_radius=0,
            scrollbar_button_color=SCROLLBAR,
            scrollbar_button_hover_color=SCROLLBAR_HOVER,
        )
        self.cards.bind("<Configure>", self._on_cards_resize)
        self.grid_holder = ctk.CTkFrame(self.cards, fg_color="transparent")
        self.grid_holder.pack(anchor="n", pady=(2, 8))

        action = ctk.CTkFrame(self.root, fg_color=CARD, corner_radius=18,
                              border_width=1, border_color=BORDER)
        action.pack(fill="x", padx=24, pady=(0, 20))
        action_inner = ctk.CTkFrame(action, fg_color="transparent")
        action_inner.pack(fill="x", padx=20, pady=13)

        info = ctk.CTkFrame(action_inner, fg_color="transparent")
        info.pack(side="left")
        self.count_label = ctk.CTkLabel(
            info, text=self.t("photo_count", count=0), font=self.f_count,
            text_color=TEXT)
        self.count_label.pack(anchor="w")
        self.hint_label = ctk.CTkLabel(
            info, text=self.t("help_hint"), font=self.f_hint, text_color=MUTED)
        self.hint_label.pack(anchor="w", pady=(1, 0))

        self.save_button = ctk.CTkButton(
            action_inner, text=self.t("save"),
            image=_icon(GLYPHS["save"], "white", 17), compound="left",
            command=self.convert, font=self.f_primary,
            fg_color=GREEN, hover_color=GREEN_HOVER,
            corner_radius=12, height=44, width=208,
        )
        self.save_button.pack(side="right")

        self.toast = ctk.CTkFrame(self.root, fg_color=ACCENT_SOFT, corner_radius=14,
                                  border_width=1, border_color="#D6E4FF")
        self.toast_label = ctk.CTkLabel(self.toast, text="", font=self.f_secondary,
                                        text_color=ACCENT, wraplength=640)
        self.toast_label.pack(padx=16, pady=7)

    def _ghost_button(self, parent, text, glyph, color, command, hover=GHOST_HOVER):
        return ctk.CTkButton(
            parent, text=text, image=_icon(glyph, color, 15), compound="left",
            command=command, font=self.f_ghost, fg_color="transparent",
            hover_color=hover, text_color=color, corner_radius=9,
            height=40, width=94,
        )

    def _build_empty(self):
        self.empty = ctk.CTkFrame(self.panel, fg_color="transparent")
        inner = ctk.CTkFrame(self.empty, fg_color="transparent")
        inner.place(relx=0.5, rely=0.46, anchor="center")

        art = ctk.CTkImage(
            light_image=_round_icon(GLYPHS["photo"], LIGHT["accent"],
                                    LIGHT["accent_soft"], 104),
            dark_image=_round_icon(GLYPHS["photo"], DARK["accent"],
                                   DARK["accent_soft"], 104),
            size=(104, 104),
        )
        ctk.CTkLabel(inner, text="", image=art).pack()

        ctk.CTkLabel(inner, text=self.t("empty_title"), font=self.f_empty_title,
                     text_color=TEXT).pack(pady=(16, 2))
        ctk.CTkLabel(
            inner, text=self.t("empty_sub"),
            font=self.f_empty_sub, text_color=MUTED,
        ).pack()

        steps = ctk.CTkFrame(inner, fg_color="transparent")
        steps.pack(pady=(22, 0))
        for number, label in (("1", self.t("empty_chip_screenshot")),
                              ("2", self.t("empty_chip_paste")),
                              ("3", self.t("empty_chip_save"))):
            chip = ctk.CTkFrame(steps, fg_color=CARD, corner_radius=12,
                                border_width=1, border_color=BORDER)
            chip.pack(side="left", padx=5)
            ctk.CTkLabel(chip, text=number, width=22, height=22, corner_radius=7,
                         fg_color=ACCENT_SOFT, text_color=ACCENT,
                         font=self.f_chip).pack(side="left", padx=(9, 7), pady=8)
            ctk.CTkLabel(chip, text=label, font=self.f_chip,
                         text_color=MUTED).pack(side="left", padx=(0, 11), pady=8)

        ctk.CTkLabel(inner, text=self.t("empty_note"),
                     font=self.f_hint, text_color=MUTED).pack(pady=(14, 0))

    def _bind_keys(self):
        self.root.bind_all("<Control-v>", self.paste)
        self.root.bind_all("<Control-V>", self.paste)
        self.root.bind_all("<Delete>", self._on_delete_key)
        self.root.bind_all("<Escape>", self._on_escape)

    # ---------- theme and language ----------

    def toggle_theme(self):
        dark = ctk.get_appearance_mode().lower() != "dark"
        self.theme = "dark" if dark else "light"
        ctk.set_appearance_mode(self.theme)
        save_theme(self.theme)
        self._update_theme_button()
        _set_dark_titlebar(self.root, dark)
        self._hide_toast()

    def _update_theme_button(self):
        dark = ctk.get_appearance_mode().lower() == "dark"
        self.theme_button.configure(
            text=self.t("theme_light") if dark else self.t("theme_dark"),
            image=self._sun_icon if dark else self._moon_icon,
        )

    def open_language_dialog(self):
        dialog = LanguageDialog(
            self.root, current=self.lang,
            localized_title=self.t("language_title"),
            localized_subtitle=self.t("language_subtitle"),
        )
        self.root.wait_window(dialog)
        if dialog.result and dialog.result != self.lang:
            self.set_language(dialog.result)

    def set_language(self, code):
        save_language(code)
        self.lang = code
        self.t = Translator(code)
        self.meta = self.t.meta
        self.font_family = CJK_FONT if code == "zh" else FONT
        self._rebuild_ui()

    def _rebuild_ui(self):
        self._cancel_toast()
        for child in self.root.winfo_children():
            child.destroy()
        self._build_fonts()
        self._build_ui()
        self._update_theme_button()
        self._build_empty()
        self._showing_empty = None
        self._render()
        self.root.title(self.t("app_title"))
        _set_dark_titlebar(self.root, self.theme == "dark")

    # ---------- events ----------

    def _on_cards_resize(self, event):
        if event.width < 150:
            return
        cols = max(2, (event.width - 8) // 196)
        if cols != self._cols:
            self._cols = cols
            if self.items:
                self._render()

    def _on_delete_key(self, _event=None):
        self.delete_selected()
        return "break"

    def _on_escape(self, _event=None):
        if self.selected is not None:
            self.selected = None
            self._render()

    # ---------- photo handling ----------

    def _add_image(self, image, name):
        image = ImageOps.exif_transpose(image)
        image.load()
        light_tile = _thumb_tile(
            image, self.THUMB, bg=LIGHT["tile_bg"], border=LIGHT["tile_border"])
        dark_tile = _thumb_tile(
            image, self.THUMB, bg=DARK["tile_bg"], border=DARK["tile_border"])
        self.items.append({
            "image": image,
            "thumb": ctk.CTkImage(light_image=light_tile, dark_image=dark_tile,
                                  size=(self.THUMB, self.THUMB)),
            "name": name,
        })
        self.selected = len(self.items) - 1

    def paste(self, _event=None):
        try:
            data = ImageGrab.grabclipboard()
        except Exception as exc:
            self._show_toast(
                self.t("toast_clipboard_error", error=exc), "error")
            return "break"

        if isinstance(data, Image.Image):
            self._add_image(data, self.t("screenshot"))
            self._render()
            self._scroll_to_end()
            self._show_toast(self.t("toast_photo_added"), "success")
        elif isinstance(data, list):
            added = 0
            skipped = 0
            for path in data:
                if not path.lower().endswith(IMAGE_EXTENSIONS):
                    skipped += 1
                    continue
                try:
                    with Image.open(path) as opened:
                        self._add_image(opened.copy(), os.path.basename(path))
                    added += 1
                except Exception:
                    skipped += 1
            if added:
                self._render()
                self._scroll_to_end()
                suffix = self.t("toast_skipped", skipped=skipped) if skipped else ""
                self._show_toast(
                    self.t("toast_files_added", added=added) + suffix,
                    "success")
            else:
                self._show_toast(self.t("toast_no_images_in_files"), "error")
        else:
            self._show_toast(self.t("toast_no_clipboard_image"), "error")

        return "break"

    def add_files(self):
        paths = filedialog.askopenfilenames(
            parent=self.root,
            title=self.t("dialog_add_images"),
            filetypes=[(self.t("filter_images"),
                        "*.png *.jpg *.jpeg *.bmp *.gif *.webp *.tif *.tiff"),
                       (self.t("filter_all_files"), "*.*")],
        )
        if not paths:
            return
        added = 0
        errors = 0
        for path in paths:
            try:
                with Image.open(path) as opened:
                    self._add_image(opened.copy(), os.path.basename(path))
                added += 1
            except Exception:
                errors += 1
        if added:
            self._render()
            self._scroll_to_end()
            suffix = self.t("toast_errors", errors=errors) if errors else ""
            self._show_toast(
                self.t("toast_files_added", added=added) + suffix,
                "success")

    def select(self, index):
        self.selected = index
        self._render()

    def delete_selected(self):
        if self.selected is None or not self.items:
            self._show_toast(self.t("toast_select_first"), "error")
            return
        index = self.selected
        self.items.pop(index)
        if self.items:
            self.selected = min(index, len(self.items) - 1)
        else:
            self.selected = None
        self._render()
        self._show_toast(self.t("toast_photo_deleted"), "info")

    def clear_all(self):
        if not self.items:
            return
        if not messagebox.askyesno(
                self.t("app_title"),
                self.t("confirm_clear", count=len(self.items)),
                parent=self.root):
            return
        self.items.clear()
        self.selected = None
        self._render()
        self._show_toast(self.t("toast_list_cleared"), "info")

    def move_selected(self, delta):
        if self.selected is None:
            self._show_toast(self.t("toast_select_first"), "error")
            return
        target = self.selected + delta
        if not 0 <= target < len(self.items):
            return
        self.items[self.selected], self.items[target] = self.items[target], self.items[self.selected]
        self.selected = target
        self._render()
        self._show_toast(self.t("toast_order_changed"), "info")

    # ---------- rendering ----------

    def _render(self):
        self.count_label.configure(
            text=self.t("photo_count", count=len(self.items)))

        if not self.items:
            if self._showing_empty is not True:
                self.cards.pack_forget()
                self.empty.pack(fill="both", expand=True, padx=6, pady=6)
                self._showing_empty = True
            return

        if self._showing_empty is not False:
            self.empty.pack_forget()
            self.cards.pack(fill="both", expand=True, padx=10, pady=10)
            self._showing_empty = False

        for child in self.grid_holder.winfo_children():
            child.destroy()

        cols = max(2, self._cols or 4)

        for index, item in enumerate(self.items):
            row, column = divmod(index, cols)
            selected = index == self.selected
            width, height = item["image"].size

            card = ctk.CTkButton(
                self.grid_holder,
                text=f"{index + 1}  ·  {width}×{height}",
                image=item["thumb"], compound="top",
                width=self.CARD_W, height=self.CARD_H,
                fg_color=SELECTED_BG if selected else CARD,
                hover_color=CARD_HOVER,
                border_width=2 if selected else 1,
                border_color=ACCENT if selected else BORDER,
                corner_radius=16,
                text_color=ACCENT if selected else MUTED,
                font=self.f_card,
                command=lambda i=index: self.select(i),
            )
            card.grid(row=row, column=column, padx=7, pady=7, sticky="n")

    # ---------- saving ----------

    def convert(self):
        if not self.items:
            messagebox.showinfo(
                self.t("app_title"), self.t("no_photos"), parent=self.root)
            return

        date = datetime.now().strftime("%Y-%m-%d_%H-%M")
        default_name = self.t("default_filename", date=date)
        path = filedialog.asksaveasfilename(
            parent=self.root,
            title=self.t("dialog_save_word"),
            defaultextension=".docx",
            initialfile=default_name,
            initialdir=desktop_path(),
            filetypes=[(self.t("filter_word"), "*.docx")],
        )
        if not path:
            return

        self.root.configure(cursor="watch")
        self.root.update_idletasks()
        try:
            save_docx([item["image"] for item in self.items], path)
        except Exception as exc:
            traceback.print_exc()
            messagebox.showerror(
                self.t("app_title"), self.t("error_save", error=exc),
                parent=self.root)
            return
        finally:
            self.root.configure(cursor="")

        if messagebox.askyesno(
                self.t("app_title"), self.t("ready_open", path=path),
                parent=self.root):
            try:
                os.startfile(path)
            except Exception as exc:
                messagebox.showerror(
                    self.t("app_title"), self.t("error_open", error=exc),
                    parent=self.root)

    def _scroll_to_end(self):
        def scroll():
            try:
                self.cards._parent_canvas.yview_moveto(1.0)
            except Exception:
                pass
        self.root.after(60, scroll)

    # ---------- notifications ----------

    def _show_toast(self, text, kind="info"):
        palette = DARK if ctk.get_appearance_mode().lower() == "dark" else LIGHT
        background, foreground, border = palette.get(
            f"toast_{kind}", palette["toast_info"])
        self.toast.configure(fg_color=background, border_color=border)
        self.toast_label.configure(text=text, text_color=foreground)
        self.toast.place(relx=0.5, rely=1.0, y=-112, anchor="s")
        self.toast.lift()

        if self._toast_job is not None:
            try:
                self.root.after_cancel(self._toast_job)
            except Exception:
                pass
        self._toast_job = self.root.after(2800, self._hide_toast)

    def _cancel_toast(self):
        if self._toast_job is not None:
            try:
                self.root.after_cancel(self._toast_job)
            except Exception:
                pass
            self._toast_job = None

    def _hide_toast(self):
        self._toast_job = None
        self.toast.place_forget()


def main():
    ctk.set_appearance_mode(load_theme())
    root = ctk.CTk()
    root.withdraw()

    language = load_language()
    if language is None:
        dialog = LanguageDialog(root)
        root.wait_window(dialog)
        language = dialog.result
        if not language:
            root.destroy()
            return
        save_language(language)

    App(root, language)
    root.deiconify()
    root.update()
    root.mainloop()


if __name__ == "__main__":
    main()
